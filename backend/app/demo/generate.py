"""Generates realistic, intentionally messy demo data for EasyETL.

Files: Customer_Data.xlsx, Orders.csv, Vehicles.json, Service_History.xml (+ a demo SQL database and a ZIP bundle).
They deliberately contain duplicates, null values, inconsistent phone/date formats, different country spellings,
invalid emails, schema inconsistencies, nested JSON and XML structures — so profiling, AI recommendations and the
Transformation Studio can demonstrate real value.

Run: python -m app.demo.generate
"""
from __future__ import annotations

import csv
import json
import random
import sqlite3
import zipfile
from datetime import date, timedelta
from pathlib import Path

DATA_DIR = Path(__file__).parent / "data"

FIRST = ["James", "Mary", "Robert", "Patricia", "John", "Jennifer", "Michael", "Linda", "David", "Elizabeth", "William",
         "Barbara", "Richard", "Susan", "Joseph", "Jessica", "Thomas", "Sarah", "Priya", "Wei", "Carlos", "Fatima",
         "Hans", "Ingrid", "Olivia", "Liam", "Noah", "Emma", "Ava", "Mateo", "Sofia", "Lukas", "Aiko", "Omar", "Chloe"]
LAST = ["Smith", "Johnson", "Williams", "Brown", "Jones", "Garcia", "Miller", "Davis", "Rodriguez", "Martinez",
        "Hernandez", "Lopez", "Wilson", "Anderson", "Taylor", "Thomas", "Moore", "Jackson", "Martin", "Lee", "Patel",
        "Müller", "Schmidt", "O'Brien", "Nguyen", "Kim", "Chen", "Singh", "Khan", "Rossi", "Dubois", "Tanaka"]
COUNTRY_VARIANTS = {
    "United States": ["USA", "United States", "US", "U.S.A.", "united states", "United States of America", " USA "],
    "United Kingdom": ["UK", "United Kingdom", "Great Britain", "U.K.", "england"],
    "Germany": ["Germany", "DE", "Deutschland", "germany"],
    "Canada": ["Canada", "CA", "CAN", "canada"],
}
COUNTRY_WEIGHTS = [("United States", 0.62), ("United Kingdom", 0.14), ("Germany", 0.12), ("Canada", 0.12)]
CITIES = {"United States": ["Austin", "Denver", "Chicago", "Seattle", "Miami", "Boston"],
          "United Kingdom": ["London", "Manchester", "Leeds"], "Germany": ["Berlin", "Munich", "Hamburg"],
          "Canada": ["Toronto", "Vancouver", "Montreal"]}
SEGMENTS = ["Retail", "Fleet", "Enterprise", "retail", "FLEET"]
MAKES = {"Toyota": ["Camry", "RAV4", "Corolla", "Highlander"], "Ford": ["F-150", "Escape", "Mustang Mach-E"],
         "Tesla": ["Model 3", "Model Y"], "BMW": ["X3", "i4", "330i"], "Volkswagen": ["ID.4", "Golf", "Tiguan"]}
SERVICE_TYPES = ["Oil Change", "Brake Service", "Tire Rotation", "Battery Replacement", "Annual Inspection", "Recall Repair"]


def _pick_country(rng: random.Random) -> str:
    r, acc = rng.random(), 0.0
    for c, w in COUNTRY_WEIGHTS:
        acc += w
        if r <= acc:
            return c
    return "United States"


def _phone(rng: random.Random, country: str) -> str | None:
    if rng.random() < 0.04:
        return None
    a, b, c = rng.randint(200, 989), rng.randint(200, 999), rng.randint(1000, 9999)
    if country == "United Kingdom":
        return rng.choice([f"+44 20 {b}{c % 10} {c}", f"020 {b}{c % 10} {c}", f"0044 20{b}{c}"])
    if country == "Germany":
        return rng.choice([f"+49 30 {b}{c}", f"030-{b}{c}", f"0049 30 {b} {c}"])
    fmt = rng.random()
    if fmt < 0.35:
        return f"({a}) {b}-{c}"
    if fmt < 0.55:
        return f"{a}.{b}.{c}"
    if fmt < 0.75:
        return f"+1 {a} {b} {c}"
    if fmt < 0.9:
        return f"{a}{b}{c}"
    return f"{a}-{b}-{c}"


def _fmt_date(rng: random.Random, d: date, clean_ratio: float = 0.7) -> str:
    if rng.random() < clean_ratio:
        return d.isoformat()
    return rng.choice([d.strftime("%m/%d/%Y"), d.strftime("%d %b %Y"), d.strftime("%Y/%m/%d"), d.strftime("%d-%m-%Y"),
                       d.strftime("%b %d, %Y")])


def _email(rng: random.Random, first: str, last: str, i: int) -> str | None:
    r = rng.random()
    if r < 0.032:
        return None
    base = f"{first}.{last}".lower().replace("'", "").replace("ü", "u")
    domain = rng.choice(["gmail.com", "outlook.com", "yahoo.com", "acme-corp.com", "mail.de", "company.co.uk"])
    if r < 0.06:
        return rng.choice([f"{base}@", f"{base}.{domain}", "n/a", f"{base}@@{domain}", "unknown", f"{base} @{domain}"])
    email = f"{base}{i % 97}@{domain}"
    return email.upper() if rng.random() < 0.05 else email


def generate_customers(rng: random.Random, n: int = 1500) -> list[dict]:
    rows = []
    for i in range(n):
        cid = f"C{10001 + i}"
        first, last = rng.choice(FIRST), rng.choice(LAST)
        country = _pick_country(rng)
        dob = date(1950, 1, 1) + timedelta(days=rng.randint(0, 365 * 52))
        signup = date(2018, 1, 1) + timedelta(days=rng.randint(0, 365 * 6))
        name_first = rng.choice([first, first.upper(), f" {first}", first.lower()]) if rng.random() < 0.12 else first
        revenue = round(rng.lognormvariate(8.5, 1.1), 2)
        if rng.random() < 0.01:
            revenue = revenue * 40  # outliers
        rows.append({
            "Customer_ID": cid,
            "First_Name": name_first,
            "Last_Name": f"{last} " if rng.random() < 0.08 else last,
            "Email": _email(rng, first, last, i),
            "Phone": _phone(rng, country),
            "Date_of_Birth": _fmt_date(rng, dob, 0.65) if rng.random() > 0.02 else None,
            "Country": rng.choice(COUNTRY_VARIANTS[country]),
            "City": rng.choice(CITIES[country]),
            "Postal_Code": str(rng.randint(10000, 99999)),
            "Segment": rng.choice(SEGMENTS) if rng.random() > 0.05 else None,
            "Signup_Date": _fmt_date(rng, signup, 0.8),
            "Annual_Revenue": (f"${revenue:,.2f}" if rng.random() < 0.3 else f"{revenue:.2f}") if rng.random() > 0.05 else None,
            "Loyalty_Points": str(rng.randint(0, 25000)),
            "Marketing_Opt_In": rng.choice(["Yes", "No", "Y", "N", "TRUE", "false", "1", "0"]),
        })
    # duplicates: exact copies + updated re-entries of the same customer
    dup_rows = []
    for _ in range(38):
        dup_rows.append(dict(rng.choice(rows)))
    for _ in range(24):
        d = dict(rng.choice(rows))
        d["Phone"] = _phone(rng, "United States")
        d["Signup_Date"] = _fmt_date(rng, date(2024, 6, 1) + timedelta(days=rng.randint(0, 90)), 1.0)
        dup_rows.append(d)
    rows.extend(dup_rows)
    rng.shuffle(rows)
    return rows


def generate_orders(rng: random.Random, customers: list[dict], vehicles: list[dict], n: int = 4200) -> list[dict]:
    ids = sorted({c["Customer_ID"] for c in customers})
    vins = [v["vin"] for v in vehicles]
    statuses = ["Completed", "completed", "COMPLETED", "Pending", "Cancelled", "Canceled", "Shipped", "shipped"]
    rows = []
    for i in range(n):
        d = date(2023, 1, 1) + timedelta(days=rng.randint(0, 600))
        amount = round(rng.uniform(25, 4800), 2)
        cust = rng.choice(ids)
        if rng.random() < 0.03:
            cust = f" {cust.lower()}"  # inconsistent code format
        if rng.random() < 0.01:
            cust = f"C{rng.randint(90000, 99999)}"  # orphan (referential integrity)
        rows.append({
            "order_id": f"ORD-{200000 + i}",
            "customer_id": cust,
            "vin": rng.choice(vins) if rng.random() > 0.2 else "",
            "order_date": _fmt_date(rng, d, 0.75),
            "product_category": rng.choice(["Vehicle", "Parts", "Service Plan", "Accessories", "Warranty"]),
            "quantity": str(rng.choice([1, 1, 1, 2, 3, 4])),
            "unit_price": (f"${amount:,.2f}" if rng.random() < 0.15 else f"{amount:.2f}") if rng.random() > 0.02 else "",
            "discount_pct": rng.choice(["0", "5", "10", "15%", "0.1", ""]),
            "currency": rng.choice(["USD", "USD", "USD", "usd", "EUR", "GBP"]),
            "status": rng.choice(statuses),
            "sales_channel": rng.choice(["Online", "Dealer", "Phone", "online"]),
        })
        if rng.random() < 0.004:
            rows[-1]["unit_price"] = f"-{amount:.2f}"  # invalid negative
    for _ in range(45):
        rows.append(dict(rng.choice(rows)))
    rng.shuffle(rows)
    return rows


def generate_vehicles(rng: random.Random, customers: list[dict], n: int = 900) -> list[dict]:
    ids = sorted({c["Customer_ID"] for c in customers})
    out = []
    for i in range(n):
        make = rng.choice(list(MAKES))
        ev = make == "Tesla" or rng.random() < 0.15
        out.append({
            "vin": f"1HGCM{rng.randint(10, 99)}{chr(65 + i % 26)}{100000 + i}",
            "make": make if rng.random() > 0.05 else make.upper(),
            "model": rng.choice(MAKES[make]),
            "year": rng.randint(2012, 2025),
            "owner": {"customer_id": rng.choice(ids), "since": _fmt_date(rng, date(2019, 1, 1) + timedelta(days=rng.randint(0, 1800)), 0.9)},
            "specs": {"engine": {"type": "Electric" if ev else rng.choice(["Gasoline", "Hybrid", "Diesel"]),
                                  "horsepower": rng.randint(140, 480)},
                      "color": rng.choice(["White", "Black", "Silver", "Blue", "Red", "grey", "Gray"]),
                      "mileage_km": rng.randint(0, 220000) if rng.random() > 0.04 else None},
            "features": rng.sample(["Navigation", "Heated Seats", "Sunroof", "Lane Assist", "Adaptive Cruise", "360 Camera"], rng.randint(0, 4)),
            "warranty": {"active": rng.random() > 0.3, "expires": _fmt_date(rng, date(2025, 1, 1) + timedelta(days=rng.randint(0, 1400)), 0.85)},
        })
    return out


def generate_service_xml(rng: random.Random, vehicles: list[dict], n: int = 2600) -> str:
    from xml.sax.saxutils import escape

    lines = ['<?xml version="1.0" encoding="UTF-8"?>', '<ServiceHistory source="DealerDMS" exported="2025-06-30">']
    for i in range(n):
        v = rng.choice(vehicles)
        d = date(2022, 1, 1) + timedelta(days=rng.randint(0, 900))
        stype = rng.choice(SERVICE_TYPES)
        cost = round(rng.uniform(40, 1900), 2)
        parts = rng.sample(["Oil Filter", "Brake Pads", "Rotor", "Battery", "Air Filter", "Wiper Blades", "Spark Plugs"], rng.randint(0, 3))
        cust_tag = "CustomerID" if rng.random() > 0.05 else "CustomerId"  # schema inconsistency
        lines.append(f'  <Service id="SRV-{500000 + i}">')
        lines.append(f"    <VIN>{v['vin']}</VIN>")
        lines.append(f"    <{cust_tag}>{v['owner']['customer_id']}</{cust_tag}>")
        lines.append(f"    <ServiceDate>{escape(_fmt_date(rng, d, 0.6))}</ServiceDate>")
        lines.append(f"    <Type>{stype}</Type>")
        lines.append(f'    <Cost currency="{rng.choice(["USD", "USD", "EUR"])}">{cost if rng.random() > 0.03 else ""}</Cost>')
        lines.append(f"    <Odometer>{rng.randint(1000, 200000)}</Odometer>")
        if parts:
            lines.append("    <Parts>")
            for p in parts:
                lines.append(f"      <Part>{p}</Part>")
            lines.append("    </Parts>")
        lines.append(f"    <Technician><Name>{rng.choice(FIRST)} {rng.choice(LAST).replace(chr(39), '')}</Name><Id>T{rng.randint(100, 160)}</Id></Technician>")
        lines.append("  </Service>")
    lines.append("</ServiceHistory>")
    return "\n".join(lines)


def write_excel(path: Path, customers: list[dict], rng: random.Random) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    headers = list(customers[0].keys())
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A68")
    for r in customers:
        ws.append([r[h] for h in headers])

    ws2 = wb.create_sheet("Dealerships")
    ws2.append(["Dealer_Code", "Dealer_Name", "Region", "Country", "Opened"])
    for i, (country, cities) in enumerate(CITIES.items()):
        for j, city in enumerate(cities):
            ws2.append([f"D{100 + i * 10 + j}", f"{city} Motors", rng.choice(["North", "South", "East", "West"]),
                        rng.choice(COUNTRY_VARIANTS[country]), _fmt_date(rng, date(2005 + j, 1 + i, 10), 0.5)])

    ws3 = wb.create_sheet("Loyalty_Tiers")
    ws3.append(["Tier", "Min_Points", "Max_Points", "Discount"])
    for t in [("Bronze", 0, 4999, "0%"), ("Silver", 5000, 9999, "5%"), ("Gold", 10000, 19999, "10%"), ("Platinum", 20000, 999999, "15%")]:
        ws3.append(list(t))

    ws4 = wb.create_sheet("Region_Targets")
    ws4.append(["Region", "Quarter", "Target_Revenue", "Owner"])
    for region in ["North", "South", "East", "West"]:
        for q in ["Q1", "Q2", "Q3", "Q4"]:
            ws4.append([region, f"2025-{q}", rng.randint(200, 900) * 1000, rng.choice(FIRST)])
    wb.save(path)


def write_demo_db(path: Path, rng: random.Random) -> None:
    """A demo relational database, accessed through the same SQLAlchemy code path as SQL Server/PostgreSQL."""
    if path.exists():
        path.unlink()
    con = sqlite3.connect(path)
    cur = con.cursor()
    cur.executescript("""
        CREATE TABLE dealers (dealer_id INTEGER PRIMARY KEY, name TEXT, city TEXT, country TEXT, opened_on TEXT, modified_at TEXT);
        CREATE TABLE sales_reps (rep_id INTEGER PRIMARY KEY, dealer_id INTEGER REFERENCES dealers(dealer_id), full_name TEXT, email TEXT, phone TEXT, hired_on TEXT, modified_at TEXT);
        CREATE TABLE inventory (stock_id INTEGER PRIMARY KEY, dealer_id INTEGER REFERENCES dealers(dealer_id), vin TEXT, make TEXT, model TEXT, model_year INTEGER, list_price REAL, status TEXT, modified_at TEXT);
    """)
    dealers = []
    for i, (country, cities) in enumerate(CITIES.items()):
        for city in cities:
            did = len(dealers) + 1
            dealers.append(did)
            cur.execute("INSERT INTO dealers VALUES (?,?,?,?,?,?)", (did, f"{city} Motors", city, country, f"20{10 + did % 12}-0{1 + did % 9}-15", "2025-06-01T08:00:00"))
    for r in range(1, 121):
        f, l = rng.choice(FIRST), rng.choice(LAST)
        cur.execute("INSERT INTO sales_reps VALUES (?,?,?,?,?,?,?)", (r, rng.choice(dealers), f"{f} {l}", f"{f}.{l}@dealers.example".lower(), _phone(rng, "United States"), f"20{15 + r % 10}-0{1 + r % 9}-01", "2025-06-01T08:00:00"))
    for s in range(1, 3001):
        make = rng.choice(list(MAKES))
        cur.execute("INSERT INTO inventory VALUES (?,?,?,?,?,?,?,?,?)", (s, rng.choice(dealers), f"5YJ{rng.randint(100000, 999999)}{s}", make, rng.choice(MAKES[make]), rng.randint(2020, 2026), round(rng.uniform(18000, 95000), 2), rng.choice(["In Stock", "Sold", "In Transit", "in stock"]), f"2025-0{1 + s % 6}-{10 + s % 18}T10:00:00"))
    con.commit()
    con.close()


def generate(out_dir: Path = DATA_DIR) -> dict[str, Path]:
    rng = random.Random(42)
    out_dir.mkdir(parents=True, exist_ok=True)
    customers = generate_customers(rng)
    vehicles = generate_vehicles(rng, customers)
    orders = generate_orders(rng, customers, vehicles)

    paths = {
        "Customer_Data.xlsx": out_dir / "Customer_Data.xlsx",
        "Orders.csv": out_dir / "Orders.csv",
        "Vehicles.json": out_dir / "Vehicles.json",
        "Service_History.xml": out_dir / "Service_History.xml",
    }
    write_excel(paths["Customer_Data.xlsx"], customers, rng)
    with paths["Orders.csv"].open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(orders[0].keys()))
        w.writeheader()
        w.writerows(orders)
    paths["Vehicles.json"].write_text(json.dumps({"vehicles": vehicles, "exported_at": "2025-06-30T12:00:00Z"}, indent=1, ensure_ascii=False))
    paths["Service_History.xml"].write_text(generate_service_xml(rng, vehicles), encoding="utf-8")

    bundle = out_dir / "Dealer_Export_Bundle.zip"
    with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as zf:
        for name in ("Orders.csv", "Vehicles.json", "Service_History.xml"):
            zf.write(paths[name], f"export/{name}")
        zf.writestr("export/README.txt", "Dealer DMS nightly export.\n")
    paths["Dealer_Export_Bundle.zip"] = bundle

    db = out_dir / "demo_dealer_db.sqlite"
    write_demo_db(db, rng)
    paths["demo_dealer_db.sqlite"] = db
    return paths


if __name__ == "__main__":
    for name, p in generate().items():
        print(f"{name:32s} {p.stat().st_size / 1024:8.1f} KB")
