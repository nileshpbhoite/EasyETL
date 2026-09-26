"""Data quality rules in Excel: a template to fill in, and an importer.

Each row is one rule and can be written three ways (first one filled wins):
  1. "SQL condition"          — a Databricks SQL predicate, validated strictly
  2. "Rule type" + parameters — not_null, unique, email, phone, date, range, in_set, regex, min_length, in_reference
  3. "Rule (plain English)"   — converted by EasyETL's rule assistant, grounded in the dataset's columns

Scoring is Red/Amber/Green: pass rate ≥ Green % → Green, ≥ Amber % → Amber, otherwise Red.
"""
from __future__ import annotations

import csv
import io
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..ai import dq_rules
from . import dq_sql
from .metadata import PipelineMetadata

HEADERS = ["Rule name", "Dataset", "Column", "Rule type", "Parameters", "Rule (plain English)", "SQL condition", "Severity", "On failure",
           "Green ≥ %", "Amber ≥ %", "Enabled"]
RULE_TYPES = ["", "not_null", "unique", "email", "phone", "date", "range", "in_set", "regex", "min_length", "in_reference"]
ON_FAILURE = {"flag & load": "flag", "flag": "flag", "warn": "flag", "quarantine": "quarantine", "park": "quarantine", "drop": "drop",
              "fail pipeline": "fail", "fail": "fail", "fail the run": "fail"}
SEVERITIES = ["Critical", "High", "Medium", "Low"]
_ALIASES = {"rule name": 0, "name": 0, "dataset": 1, "table": 1, "sheet": 1, "column": 2, "field": 2, "rule type": 3, "type": 3, "check": 3,
            "parameters": 4, "params": 4, "rule (plain english)": 5, "rule": 5, "description": 5, "business rule": 5, "plain english": 5,
            "sql condition": 6, "sql": 6, "condition": 6, "expression": 6, "severity": 7, "priority": 7, "on failure": 8, "action": 8,
            "on fail": 8, "handling": 8, "green ≥ %": 9, "green": 9, "green %": 9, "green threshold": 9, "amber ≥ %": 10, "amber": 10,
            "amber %": 10, "amber threshold": 10, "enabled": 11, "active": 11}


def _short(name: str) -> str:
    return name.split(" › ")[-1]


def build_template(meta: PipelineMetadata, columns_by_dataset: dict[str, list[str]], types_by_dataset: dict[str, dict[str, str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rules"
    head_fill = PatternFill("solid", fgColor="4655EC")
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    widths = [28, 22, 20, 14, 22, 46, 40, 11, 16, 10, 10, 9]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w

    # Example rows grounded in the real columns of this pipeline
    datasets = [d for d in meta.selected_datasets() if d.id in columns_by_dataset]
    examples: list[list[Any]] = []
    if datasets:
        d = datasets[0]
        cols, types = columns_by_dataset[d.id], types_by_dataset.get(d.id, {})
        email = next((c for c in cols if types.get(c) == "email" or "mail" in c.lower()), None)
        date = next((c for c in cols if types.get(c) in ("date", "date_of_birth", "timestamp")), None)
        key = next((c for c in cols if re.search(r"(^|_)id$", c, re.I)), cols[0] if cols else None)
        num = next((c for c in cols if types.get(c) in ("integer", "decimal", "currency")), None)
        if key:
            examples.append([f"{key} is unique", _short(d.name), key, "unique", "", "", "", "Critical", "Quarantine", 100, 99, "Yes"])
        if email:
            examples.append([f"{email} is a valid email", _short(d.name), email, "", "", f"{email} must be a valid email address", "", "High", "Flag & load", 99, 95, "Yes"])
        if date:
            examples.append([f"{date} not in the future", _short(d.name), date, "", "", f"{date} cannot be in the future", "", "Medium", "Flag & load", 99.5, 97, "Yes"])
        if num:
            examples.append([f"{num} is not negative", _short(d.name), "", "", "", "", f"{num} >= 0", "Medium", "Quarantine", 99, 95, "Yes"])
        if len(cols) > 3:
            examples.append([f"{cols[1]} is required", _short(d.name), cols[1], "not_null", "", "", "", "High", "Flag & load", 98, 90, "Yes"])
    for row in examples:
        ws.append(row)

    n = 500
    def dv(values: list[str], col: str) -> None:
        v = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        ws.add_data_validation(v)
        v.add(f"{col}2:{col}{n}")
    if datasets and len(",".join(_short(d.name) for d in datasets)) < 240:
        dv([_short(d.name) for d in datasets], "B")
    dv([t for t in RULE_TYPES if t], "D")
    dv(SEVERITIES, "H")
    dv(["Flag & load", "Quarantine", "Drop", "Fail pipeline"], "I")
    dv(["Yes", "No"], "L")
    ws.freeze_panes = "A2"

    ref = wb.create_sheet("Columns")
    ref.append(["Dataset", "Column", "Detected type"])
    for cell in ref[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
    for d in datasets:
        for c in columns_by_dataset[d.id]:
            ref.append([_short(d.name), c, types_by_dataset.get(d.id, {}).get(c, "")])
    for col, w in zip("ABC", (26, 30, 18)):
        ref.column_dimensions[col].width = w

    how = wb.create_sheet("How to")
    lines = [
        "EasyETL data quality rules — one rule per row on the 'Rules' sheet.",
        "",
        "Write each rule in ONE of three ways (EasyETL uses the first one filled in):",
        "  1. SQL condition — a Databricks SQL condition that is TRUE for a good record, e.g.  amount > 0 AND currency IN ('USD','EUR')",
        "  2. Rule type + Column (+ Parameters) — not_null, unique, email, phone, date, range, in_set, regex, min_length, in_reference",
        "       Parameters: range → min=0; max=120   in_set → values=Open,Closed   regex → pattern=^[0-9]{5}$   min_length → length=3",
        "  3. Rule (plain English) — e.g. 'Ship date must be after order date', 'If status is Closed then close date is required'",
        "",
        "Severity (Critical/High/Medium/Low) weights the rule in the overall DQ score.",
        "On failure: Flag & load = record is loaded with _dq_issues/_dq_status columns;  Quarantine = record is parked in the DQ quarantine table",
        "            instead of the target;  Drop = record is discarded;  Fail pipeline = the run stops.",
        "RAG scoring: pass rate ≥ Green % → Green, ≥ Amber % → Amber, otherwise Red. Defaults: Green 99, Amber 95.",
        "Column names are listed on the 'Columns' sheet. Upload the file on the pipeline's Data Quality tab — you'll see a preview before anything is added.",
    ]
    for line in lines:
        how.append([line])
    how.column_dimensions["A"].width = 140
    how["A1"].font = Font(bold=True, size=13)
    wb.move_sheet("How to", offset=-2)
    wb.active = 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _rows(filename: str, content: bytes) -> list[list[Any]]:
    if filename.lower().endswith((".csv", ".txt")):
        text = content.decode("utf-8-sig", errors="replace")
        return [r for r in csv.reader(io.StringIO(text))]
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb["Rules"] if "Rules" in wb.sheetnames else next(s for s in wb.worksheets if s.title != "How to")
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _params(rule: str, raw: str) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for part in re.split(r"[;\n]", raw or ""):
        if "=" not in part:
            continue
        k, v = part.split("=", 1)
        k, v = k.strip().lower(), v.strip()
        if k in ("values", "value", "list"):
            out["values"] = [x.strip() for x in v.split(",") if x.strip()]
        elif k in ("min", "max"):
            out[k] = float(v) if re.fullmatch(r"-?\d+(\.\d+)?", v) else v
        elif k == "length":
            out["length"] = int(float(v))
        elif k == "pattern":
            out["pattern"] = v
    if rule == "in_set" and "values" not in out and raw:
        out["values"] = [x.strip() for x in raw.split(",") if x.strip()]
    return out


def _num(v: Any, default: float) -> float:
    if v in (None, ""):
        return default
    try:
        return float(str(v).strip().rstrip("%"))
    except ValueError:
        raise ValueError(f"'{v}' isn't a number")


def parse_rules(filename: str, content: bytes, meta: PipelineMetadata, columns_by_dataset: dict[str, list[str]],
                types_by_dataset: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    """Every row → {"row", "ok", "error" | "rule": QualityRule-compatible dict, "summary"}. Nothing is saved here."""
    rows = _rows(filename, content)
    header_idx = next((i for i, r in enumerate(rows[:10]) if sum(1 for c in r if c and str(c).strip().lower() in _ALIASES) >= 2), None)
    if header_idx is None:
        raise ValueError("We couldn't find the header row. Use the EasyETL template (columns such as 'Column', 'Rule type', 'SQL condition').")
    mapping = {i: _ALIASES[str(c).strip().lower()] for i, c in enumerate(rows[header_idx]) if c and str(c).strip().lower() in _ALIASES}
    datasets = [d for d in meta.selected_datasets() if d.id in columns_by_dataset]
    by_name = {}
    for d in datasets:
        by_name[d.name.lower()] = d
        by_name[_short(d.name).lower()] = d
    out = []
    for n, raw in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        vals: list[Any] = [None] * 12
        for i, slot in mapping.items():
            if i < len(raw):
                vals[slot] = raw[i]
        v = [("" if x is None else str(x).strip()) for x in vals]
        name, ds_name, column, rtype, params, plain, sql, sev, onfail, green, amber, enabled = v
        if not any([column, rtype, plain, sql]):
            continue
        try:
            ds = by_name.get(ds_name.lower()) if ds_name else None
            if ds_name and not ds:
                raise ValueError(f"Unknown dataset '{ds_name}'. Use one of: {', '.join(_short(d.name) for d in datasets)}")
            if not ds:
                mentioned = column or sql or plain
                ds = next((d for d in datasets if any(c.lower() in mentioned.lower() for c in columns_by_dataset[d.id])), datasets[0] if datasets else None)
            if not ds:
                raise ValueError("This pipeline has no analyzed datasets yet.")
            cols, types = columns_by_dataset[ds.id], types_by_dataset.get(ds.id, {})
            real_col = None
            if column:
                real_col = next((c for c in cols if c.lower() == column.lower() or re.sub(r"[\s_]+", "", c.lower()) == re.sub(r"[\s_]+", "", column.lower())), None)
                if not real_col:
                    raise ValueError(f"No column '{column}' in {_short(ds.name)}")
            rule: dict[str, Any] = {"dataset_id": ds.id, "origin": "excel"}
            if sql:
                norm, used = dq_sql.validate(sql, cols)
                rule.update(rule="expression", params={"sql": norm}, column=used[0] if len(used) == 1 else None, dimension="custom")
                how = "SQL"
            elif rtype:
                rt = rtype.lower().replace(" ", "_")
                if rt not in RULE_TYPES or not rt:
                    raise ValueError(f"Unknown rule type '{rtype}'")
                if not real_col:
                    raise ValueError(f"'{rtype}' needs a Column")
                from .quality import RULE_CATALOG

                rule.update(rule=rt, column=real_col, params=_params(rt, params),
                            dimension=next((c["dimension"] for c in RULE_CATALOG if c["rule"] == rt), "validity"))
                how = "Rule type"
            else:
                text = plain if not real_col or real_col.lower() in plain.lower() else f"{real_col} {plain}"
                draft = dq_rules.convert(text, cols, types)
                rule.update(rule=draft["rule"], params=draft.get("params") or {}, column=draft.get("column"), dimension=draft.get("dimension", "custom"))
                rule["source_text"] = plain
                how = "Plain English"
            action = ON_FAILURE.get(onfail.lower()) if onfail else None
            if onfail and not action:
                raise ValueError(f"Unknown 'On failure' value '{onfail}' (use Flag & load, Quarantine, Drop or Fail pipeline)")
            if action:
                rule["on_fail"] = action
            if sev:
                if sev.capitalize() not in SEVERITIES:
                    raise ValueError(f"Unknown severity '{sev}'")
                rule["severity"] = sev.lower()
            g, a = _num(green, 99.0), _num(amber, 95.0)
            if not 0 <= a <= g <= 100:
                raise ValueError("Thresholds must satisfy 0 ≤ Amber ≤ Green ≤ 100")
            rule.update(threshold_green=g, threshold_amber=a, enabled=enabled.lower() not in ("no", "false", "0", "n"))
            rule["name"] = name or plain or (f"{real_col}: {rule['rule'].replace('_', ' ')}" if real_col else (rule["params"].get("sql") or "")[:80])
            rule["description"] = rule["name"]
            out.append({"row": n, "ok": True, "rule": rule, "how": how, "dataset": _short(ds.name),
                        "summary": rule["params"].get("sql") or f"{rule.get('column')}: {rule['rule'].replace('_', ' ')}"})
        except (ValueError, dq_sql.DQSqlError) as e:
            out.append({"row": n, "ok": False, "error": str(e), "input": plain or sql or f"{column} {rtype}".strip()})
    return out
