"""Data quality authoring (plain English, SQL, Excel), RAG scoring, record handling, targets and Databricks connections."""
import io
import json
import sys
from pathlib import Path

import polars as pl
from openpyxl import load_workbook

from tests.conftest import auth


import pytest


@pytest.fixture(autouse=True)
def _cleanup(client):
    """These tests reuse the demo file; remove their pipelines so table names don't collide with other tests."""
    yield
    h = auth(client)
    for p in client.get("/api/pipelines", headers=h).json():
        if p["name"] == "DQ pipeline":
            client.delete(f"/api/pipelines/{p['id']}", headers=h)


def _pipeline(client, h):
    pid = client.post("/api/pipelines", json={"name": "DQ pipeline"}, headers=h).json()["id"]
    f = client.post("/api/demo-files/Customer_Data.xlsx/import", headers=h).json()
    r = client.post(f"/api/pipelines/{pid}/source", json={"connector": "file_upload", "file_ids": [f["id"]]}, headers=h).json()
    cust = next(d for d in r["metadata"]["source"]["datasets"] if d["name"].endswith("Customers"))
    client.post(f"/api/pipelines/{pid}/analyze", headers=h)
    return pid, cust["id"]


def test_plain_english_rules_and_scoring(client):
    h = auth(client)
    pid, ds = _pipeline(client, h)

    d = client.post(f"/api/pipelines/{pid}/quality-rules/draft", json={"dataset_id": ds, "text": "Email must be a valid email address, quarantine bad ones"}, headers=h)
    assert d.status_code == 200, d.text
    d = d.json()
    assert d["rule"]["rule"] == "email" and d["rule"]["on_fail"] == "quarantine"
    assert d["preview"]["failed"] > 0 and d["preview"]["failing_rows"] and "RLIKE" in d["sql"]
    assert "F.expr" in d["python"]

    d2 = client.post(f"/api/pipelines/{pid}/quality-rules/draft", json={"dataset_id": ds, "text": "Signup date cannot be in the future"}, headers=h).json()
    assert d2["rule"]["rule"] == "expression" and "CURRENT_DATE" in d2["sql"].upper()

    bad = client.post(f"/api/pipelines/{pid}/quality-rules/draft", json={"dataset_id": ds, "text": "Emial_Adress IS NOT NULL"}, headers=h)
    assert bad.status_code == 400

    for draft in (d, d2):
        r = client.post(f"/api/pipelines/{pid}/quality-rules", json=draft["rule"], headers=h)
        assert r.status_code == 200, r.text
    # SQL rules are validated strictly: statements are rejected
    r = client.post(f"/api/pipelines/{pid}/quality-rules", json={"dataset_id": ds, "rule": "expression", "params": {"sql": "1=1; DROP TABLE x"}}, headers=h)
    assert r.status_code == 400

    q = client.get(f"/api/pipelines/{pid}/quality", headers=h).json()
    res = next(x for x in q["datasets"] if x["dataset_id"] == ds)["after"]
    email = next(x for x in res["rules"] if x["name"] == d["rule"]["name"])
    assert email["rag"] in ("green", "amber", "red") and email["failed"] > 0
    assert res["handling"]["quarantined"] >= email["failed"] and res["handling"]["loaded"] == res["records"] - res["handling"]["quarantined"] - res["handling"]["dropped"]

    fails = client.get(f"/api/pipelines/{pid}/quality-rules/{email['rule_id']}/failures", headers=h).json()
    assert fails["total"] == email["failed"] and fails["rows"]

    # switching the pipeline default to "flag" and applying it to all rules
    r = client.put(f"/api/pipelines/{pid}/quality-config", json={"default_action": "flag", "apply_to_all": True, "fail_run_below": 80}, headers=h).json()
    assert all(x["on_fail"] == "flag" for x in r["metadata"]["quality_rules"] if x["on_fail"] not in ("drop", "fail"))
    spec = client.get(f"/api/pipelines/{pid}/spec", headers=h).json()
    assert spec["quality"]["fail_run_below"] == 80 and all(x["sql"] or x["rule"] in ("unique", "in_dataset") for x in spec["quality_rules"])


def test_excel_template_and_import(client):
    h = auth(client)
    pid, ds = _pipeline(client, h)
    t = client.get(f"/api/pipelines/{pid}/quality-rules/template", headers=h)
    assert t.status_code == 200
    wb = load_workbook(io.BytesIO(t.content))
    assert {"How to", "Rules", "Columns"} <= set(wb.sheetnames)
    ws = wb["Rules"]
    examples = ws.max_row - 1
    assert examples >= 3
    ws.append(["Revenue positive", "", "", "", "", "", "Annual_Revenue >= 0", "High", "Quarantine", 99, 90, "Yes"])
    ws.append(["Bad column", "", "Nope", "not_null", "", "", "", "", "", "", "", ""])
    ws.append(["Status set", "", "", "", "", "If country is US then phone is required", "", "Low", "Flag & load", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    files = {"file": ("rules.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    dry = client.post(f"/api/pipelines/{pid}/quality-rules/import", files=files, data={"dry_run": "true"}, headers=h).json()
    assert dry["dry_run"] and dry["invalid"] == 1 and dry["valid"] == examples + 2, dry
    assert any("Nope" in (r.get("error") or "") for r in dry["rows"])
    done = client.post(f"/api/pipelines/{pid}/quality-rules/import", files=files, data={"dry_run": "false"}, headers=h).json()
    assert not done["dry_run"] and len([r for r in done["metadata"]["quality_rules"] if r["origin"] == "excel"]) == examples + 2


def test_targets_and_connections(client):
    h = auth(client)
    pid, _ = _pipeline(client, h)
    # a source-only connector can't be saved as a target
    r = client.post("/api/connections", json={"connector": "workday", "usage": "target", "config": {"environment": "sandbox"}}, headers=h)
    assert r.status_code == 400
    sql = client.post("/api/connections", json={"connector": "sqlserver", "name": "Reporting DB", "usage": "target", "auth_method": "password",
                                                "config": {"host": "demo", "database": "reporting"}, "secrets": {"username": "etl", "password": "s3cret"}}, headers=h).json()
    assert sql["usage"] == "target" and sql["status"] == "connected"
    assert any(c["id"] == sql["id"] for c in client.get("/api/connections?usage=target", headers=h).json())
    assert not any(c["id"] == sql["id"] for c in client.get("/api/connections?usage=source", headers=h).json())

    meta = client.get(f"/api/pipelines/{pid}", headers=h).json()["metadata"]
    gold = next(t["name"] for t in meta["lakehouse"]["tables"] if t["layer"] == "gold")
    r = client.put(f"/api/pipelines/{pid}/targets", json=[{"connection_id": sql["id"], "mode": "merge", "merge_keys": [], "tables": [gold]}], headers=h)
    assert r.status_code == 400  # merge needs keys
    r = client.put(f"/api/pipelines/{pid}/targets", json=[{"connection_id": sql["id"], "mode": "merge", "merge_keys": ["customer_id"], "tables": [gold], "destination": "dbo"}], headers=h)
    assert r.status_code == 200, r.text
    files = client.get(f"/api/pipelines/{pid}/bundle", headers=h).json()["files"]
    spec = json.loads(files["config/pipeline_spec.json"])
    assert spec["targets"][0]["config"]["database"] == "reporting"
    assert "s3cret" not in json.dumps(files)  # secrets never go into the bundle
    assert "publish_targets" in next(v for k, v in files.items() if k.endswith(".job.yml"))

    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "app" / "deploy" / "runtime"))
    import easyetl_targets

    plan = easyetl_targets.publish_plan(spec)
    assert plan and plan[0]["source"].endswith(f".{gold}")
    assert easyetl_targets.merge_sql("sqlserver", "dbo.x", "dbo.x__easyetl_stage", ["id"], ["id", "v"]).startswith("MERGE INTO dbo.x AS t")


def test_databricks_connection_auth_methods(client):
    h = auth(client)
    spec = next(c for cat in client.get("/api/connectors", headers=h).json()["categories"] for c in cat["connectors"] if c["id"] == "databricks")
    assert {m["id"] for m in spec["auth_methods"]} == {"pat", "oauth_m2m", "azure_sp", "azure_msi", "azure_cli", "gcp_sa", "cli_profile"}
    assert set(spec["roles"]) == {"source", "target"}
    r = client.post("/api/connections/test", json={"connector": "databricks", "auth_method": "oauth_m2m", "config": {"host": "demo", "client_id": "abc"},
                                                   "secrets": {"client_secret": "x"}}, headers=h).json()
    assert r["ok"] and r["info"]["Authentication"] == "OAuth service principal (M2M)"
    # a real host with missing credentials fails with a clear message (no network call needed)
    r = client.post("/api/connections/test", json={"connector": "databricks", "auth_method": "pat", "config": {"host": "https://adb-1.azuredatabricks.net"}}, headers=h).json()
    assert not r["ok"] and "token" in r["message"].lower()


def test_dq_sql_engine():
    from app.engine.dq_sql import DQSqlError, evaluate, validate

    df = pl.DataFrame({"Age": ["5", "40", None], "Start": ["2024-01-01", "2024-02-01", "x"], "End": ["2024-01-05", "2024-01-01", None]})
    sql, used = validate("age between 18 and 120 and to_date(end) >= to_date(start)", df.columns)
    assert used == ["Age", "End", "Start"]
    assert evaluate(df, sql).to_list() == [False, False, True]  # NULLs pass; use not-null rules for missing values
    for bad in ("SELECT 1", "Age > 1; DELETE FROM t", "Age IN (SELECT 1)", "evil(Age)", "Agee > 1"):
        try:
            validate(bad, df.columns)
            raise AssertionError(bad)
        except DQSqlError:
            pass


def test_sql_looking_input_is_not_reinterpreted():
    from app.ai.dq_rules import convert
    from app.engine.dq_sql import DQSqlError

    cols = ["Date_of_Birth", "Annual_Revenue"]
    try:
        convert("date_of_birth <= current_date() AND age >= 18", cols)
        raise AssertionError("should report the unknown column")
    except DQSqlError as e:
        assert "age" in str(e)
    assert convert("annual revenue must not be negative", cols)["params"]["sql"] == "Annual_Revenue >= 0"
